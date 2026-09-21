import os, re, sqlite3, csv, json, threading, time, traceback
from io import BytesIO, StringIO
from flask import Flask, request, jsonify
import requests

app = Flask(__name__)
_local = threading.local()

def http():
    if not hasattr(_local, 'session'):
        _local.session = requests.Session()
        _local.session.headers.update({'User-Agent':'Zyro/4.0'})
    return _local.session

BOT_TOKEN = os.getenv('BOT_TOKEN','').strip()
ZENTO_API_BASE = os.getenv('ZENTO_API_BASE','https://zento.up.railway.app/api/bot').rstrip('/')
DB_PATH = os.getenv('DB_PATH','bot.db')
OPENROUTER_API_KEY = os.getenv('OPENROUTER_API_KEY','').strip()
OPENROUTER_MODEL = os.getenv('OPENROUTER_MODEL','openai/gpt-oss-20b').strip()
OPENROUTER_TIMEOUT = int(os.getenv('AI_TIMEOUT','30'))
HF_TOKEN = os.getenv('HF_TOKEN','').strip()
HF_IMAGE_PROVIDER = os.getenv('HF_IMAGE_PROVIDER','auto').strip() or 'auto'
HF_IMAGE_MODELS = [x.strip() for x in os.getenv('HF_IMAGE_MODELS','black-forest-labs/FLUX.1-Krea-dev,Qwen/Qwen-Image,black-forest-labs/FLUX.1-dev,black-forest-labs/FLUX.1-schnell').split(',') if x.strip()]
AI_SYSTEM_PROMPT = os.getenv('AI_SYSTEM_PROMPT','You are Zyro, a fast, smart, friendly Persian-first AI assistant. Answer in the user language. Be concise by default, structured when useful, and never claim to have done an action you could not perform.').strip()
MAX_HISTORY = int(os.getenv('MAX_HISTORY','8'))

# ---------- database ----------
def db():
    c=sqlite3.connect(DB_PATH, timeout=8)
    c.row_factory=sqlite3.Row
    return c

def init_db():
    c=db()
    c.execute('PRAGMA journal_mode=WAL')
    c.execute('CREATE TABLE IF NOT EXISTS history(user_id TEXT NOT NULL,role TEXT NOT NULL,content TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP)')
    c.execute("CREATE TABLE IF NOT EXISTS sessions(user_id TEXT PRIMARY KEY,mode TEXT DEFAULT 'chat')")
    c.commit(); c.close()

def set_mode(uid, mode):
    c=db(); c.execute("INSERT INTO sessions(user_id,mode) VALUES(?,?) ON CONFLICT(user_id) DO UPDATE SET mode=excluded.mode",(str(uid),mode)); c.commit(); c.close()

def get_mode(uid):
    c=db(); r=c.execute('SELECT mode FROM sessions WHERE user_id=?',(str(uid),)).fetchone(); c.close(); return r['mode'] if r else 'chat'

def clear_history(uid):
    c=db(); c.execute('DELETE FROM history WHERE user_id=?',(str(uid),)); c.commit(); c.close()

def get_history(uid):
    c=db(); rows=c.execute('SELECT role,content FROM history WHERE user_id=? ORDER BY rowid DESC LIMIT ?',(str(uid),MAX_HISTORY)).fetchall(); c.close()
    return [{'role':r['role'],'content':r['content']} for r in reversed(rows)]

def save_history(uid, role, content):
    c=db(); c.execute('INSERT INTO history(user_id,role,content) VALUES(?,?,?)',(str(uid),role,str(content)))
    c.execute('DELETE FROM history WHERE user_id=? AND rowid NOT IN (SELECT rowid FROM history WHERE user_id=? ORDER BY rowid DESC LIMIT ?)',(str(uid),str(uid),MAX_HISTORY))
    c.commit(); c.close()

# ---------- Telegram-ish API ----------
def api_call(method,payload=None,timeout=8):
    if not BOT_TOKEN: return False,None
    try:
        r=http().post(f'{ZENTO_API_BASE}/{BOT_TOKEN}/{method}',json=payload or {},timeout=timeout)
        try: data=r.json()
        except Exception: data=None
        return r.ok,data
    except requests.RequestException: return False,None

def send_message(cid,text,reply_markup=None):
    p={'chat_id':str(cid),'text':str(text)}
    if reply_markup: p['reply_markup']=reply_markup
    return api_call('sendMessage',p,8)[0]

def answer_callback(qid):
    if qid: api_call('answerCallbackQuery',{'callback_query_id':qid},5)

def send_photo(cid,data,caption=''):
    try:
        form={'chat_id':str(cid)}
        if caption: form['caption']=caption
        r=http().post(f'{ZENTO_API_BASE}/{BOT_TOKEN}/sendPhoto',data=form,files={'photo':('zyro.png',data,'image/png')},timeout=45)
        return r.ok
    except requests.RequestException: return False

def send_document(cid,data,filename='zyro.txt',caption=''):
    try:
        form={'chat_id':str(cid)}
        if caption: form['caption']=caption
        mime='application/octet-stream'
        if filename.endswith('.txt'): mime='text/plain'
        elif filename.endswith('.json'): mime='application/json'
        elif filename.endswith('.csv'): mime='text/csv'
        elif filename.endswith('.pdf'): mime='application/pdf'
        elif filename.endswith('.docx'): mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        elif filename.endswith('.xlsx'): mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        r=http().post(f'{ZENTO_API_BASE}/{BOT_TOKEN}/sendDocument',data=form,files={'document':(filename,data,mime)},timeout=45)
        return r.ok
    except requests.RequestException: return False

def btn(text,data): return {'text':text,'callback_data':data}
def menu(rows): return {'inline_keyboard':rows}

# ---------- UI ----------
def home(cid):
    text=('⚡️ ZYRO AI 4.0\n\n'
          'دستیار هوشمند سریع برای گفتگو، تصویر و ساخت فایل.\n\n'
          'از دکمه‌های زیر یک قابلیت را انتخاب کن:')
    send_message(cid,text,menu([
        [btn('💬 گفت‌وگوی هوشمند','chat'),btn('🖼️ ساخت تصویر','image')],
        [btn('📄 ساخت فایل','file'),btn('🧠 ابزارهای بیشتر','tools')],
        [btn('🧹 گفت‌وگوی جدید','clear'),btn('ℹ️ راهنما','help')]
    ]))

def tools_menu(cid):
    send_message(cid,'🧠 ابزارهای Zyro\n\nهرکدام را انتخاب کن و بعد درخواستت را بفرست:',menu([
        [btn('📝 خلاصه‌سازی','summary'),btn('🌍 ترجمه','translate')],
        [btn('💻 کدنویسی','code'),btn('📊 گزارش','report')],
        [btn('📄 ساخت فایل','file'),btn('🏠 خانه','home')]
    ]))

def help_text(cid):
    send_message(cid,'ℹ️ راهنمای حرفه‌ای Zyro\n\n'
                 '💬 گفت‌وگو — چت با هوش مصنوعی\n'
                 '🖼️ ساخت تصویر — توضیح تصویر را بفرست\n'
                 '📄 ساخت فایل — TXT / MD / CSV / JSON / PDF / DOCX / XLSX\n'
                 '🧠 ابزارها — خلاصه، ترجمه، کدنویسی و گزارش\n'
                 '👥 گروه — Zyro یا @Zyro را صدا بزن\n'
                 '🧹 گفت‌وگوی جدید — حافظه همین کاربر پاک می‌شود',menu([[btn('🏠 خانه','home')]]))

# ---------- AI ----------
def chat(prompt,uid,system_extra=''):
    if not OPENROUTER_API_KEY:
        return None,'کلید OPENROUTER_API_KEY در Railway تنظیم نشده است.'
    headers={'Authorization':f'Bearer {OPENROUTER_API_KEY}','Content-Type':'application/json','X-Title':'Zyro AI 4.0'}
    system=AI_SYSTEM_PROMPT + (('\n'+system_extra) if system_extra else '')
    payload={'model':OPENROUTER_MODEL,'messages':[{'role':'system','content':system}]+get_history(uid)+[{'role':'user','content':prompt}], 'temperature':0.25,'max_completion_tokens':700}
    try:
        r=http().post('https://openrouter.ai/api/v1/chat/completions',headers=headers,json=payload,timeout=OPENROUTER_TIMEOUT)
        if not r.ok:
            if r.status_code==404: return None,'مدل گفتگو پیدا نشد. مقدار OPENROUTER_MODEL را بررسی کن.'
            if r.status_code in (401,403): return None,'کلید OpenRouter معتبر نیست یا دسترسی کافی ندارد.'
            if r.status_code==429: return None,'سرویس شلوغ است؛ چند لحظه بعد دوباره تلاش کن.'
            return None,f'سرویس گفتگو خطا داد ({r.status_code}).'
        data=r.json(); answer=((data.get('choices') or [{}])[0].get('message') or {}).get('content')
        if not answer: return None,'پاسخ خالی دریافت شد.'
        answer=str(answer).strip(); save_history(uid,'user',prompt); save_history(uid,'assistant',answer); return answer,None
    except requests.Timeout: return None,'زمان پاسخ‌گویی تمام شد؛ دوباره امتحان کن.'
    except requests.RequestException: return None,'ارتباط با سرویس گفتگو برقرار نشد.'
    except Exception: return None,'خطای داخلی در پاسخ‌گویی.'

def image(prompt):
    if not HF_TOKEN: return None,'HF_TOKEN در Railway تنظیم نشده است.'
    try:
        from huggingface_hub import InferenceClient
    except ImportError: return None,'huggingface_hub نصب نشده است.'
    last=[]
    # One client per request + auto provider: HF can select a fast available provider and fail over.
    for model in HF_IMAGE_MODELS:
        try:
            client=InferenceClient(provider=HF_IMAGE_PROVIDER,api_key=HF_TOKEN,timeout=75)
            pic=client.text_to_image(prompt,model=model)
            b=BytesIO(); pic.save(b,format='PNG'); return b.getvalue(),None
        except Exception as e:
            last.append(f'{model}: {str(e)[:160]}')
    joined=' | '.join(last).lower()
    if '402' in joined or 'payment' in joined or 'credit' in joined:
        return None,'اعتبار ساخت تصویر کافی نیست؛ Billing/Inference Providers حساب Hugging Face را بررسی کن.'
    if '401' in joined or '403' in joined:
        return None,'HF_TOKEN دسترسی Inference Providers ندارد.'
    return None,'هیچ مدل تصویر در دسترس نبود. /health را بررسی کن.'

# ---------- files ----------
def make_file(prompt,answer):
    low=prompt.lower()
    if any(x in low for x in ('xlsx','اکسل','excel')):
        try:
            from openpyxl import Workbook
            wb=Workbook(); ws=wb.active; ws.title='Zyro'
            for i,line in enumerate(answer.splitlines(),1): ws.cell(i,1,line)
            b=BytesIO(); wb.save(b); return b.getvalue(),'zyro.xlsx'
        except Exception: pass
    if any(x in low for x in ('docx','word','ورد')):
        try:
            from docx import Document
            d=Document();
            for line in answer.splitlines(): d.add_paragraph(line)
            b=BytesIO(); d.save(b); return b.getvalue(),'zyro.docx'
        except Exception: pass
    if any(x in low for x in ('pdf','پی‌دی‌اف','پی دی اف')):
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.pdfbase import pdfmetrics
            b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); c.setFont('Helvetica',11)
            y=800
            for line in answer.splitlines():
                if y<40: c.showPage(); c.setFont('Helvetica',11); y=800
                c.drawString(40,y,line[:110]); y-=16
            c.save(); return b.getvalue(),'zyro.pdf'
        except Exception: pass
    if any(x in low for x in ('json','جیسون')):
        obj={'zyro':True,'content':answer}
        return json.dumps(obj,ensure_ascii=False,indent=2).encode('utf-8'),'zyro.json'
    if any(x in low for x in ('csv','اکسل','excel')):
        out=StringIO(); w=csv.writer(out); w.writerow(['Zyro'])
        for line in answer.splitlines(): w.writerow([line])
        return out.getvalue().encode('utf-8-sig'),'zyro.csv'
    if any(x in low for x in ('markdown','مارک‌داون','مارک داون','.md')):
        return answer.encode('utf-8'),'zyro.md'
    return answer.encode('utf-8'),'zyro.txt'

# ---------- parsing ----------
def is_mention(text): return bool(text and re.search(r'(^|[^a-z0-9_])@?zyro([^a-z0-9_]|$)',str(text),re.I))
def strip_mention(text): return re.sub(r'\s+',' ',re.sub(r'(^|[^a-z0-9_])@?zyro(?=$|[^a-z0-9_])',' ',str(text),flags=re.I)).strip(' ,:؛،')

def parse_message(data):
    m=data.get('message') or data.get('msg') or data.get('channel_post')
    if not isinstance(m,dict): return None
    ch=m.get('chat') or {}; sender=m.get('from') or m.get('sender') or {}
    return {'chat_id':ch.get('id'),'chat_type':ch.get('type',''),'user_id':sender.get('id') or ch.get('id'),'text':(m.get('text') or '').strip(),'message':m}

def parse_callback(data):
    q=data.get('callback_query') or data.get('callback')
    if not isinstance(q,dict): return None
    m=q.get('message') or q.get('msg') or {}; ch=m.get('chat') or {}; s=q.get('from') or q.get('sender') or {}
    return ch.get('id'),s.get('id') or ch.get('id'),str(q.get('data') or q.get('callback_data') or ''),q.get('id') or q.get('callback_query_id')

# ---------- handlers ----------
def handle_text(cid,uid,text,extra=''):
    answer,err=chat(text,uid,extra)
    if err: send_message(cid,'❌ '+err,menu([[btn('🏠 خانه','home'),btn('🔄 دوباره','chat')]])); return
    send_message(cid,'✨ Zyro\n\n'+answer,menu([[btn('💬 ادامه گفت‌وگو','chat'),btn('🖼️ ساخت تصویر','image')],[btn('📄 تبدیل به فایل','file'),btn('🏠 خانه','home')]]))

def handle_image(cid,prompt):
    data,err=image(prompt)
    if err: send_message(cid,'❌ '+err,menu([[btn('🔄 تلاش دوباره','image'),btn('🏠 خانه','home')]])); return
    if not send_photo(cid,data,'✨ ساخته‌شده با Zyro AI'):
        send_message(cid,'❌ تصویر ساخته شد اما ارسال آن ناموفق بود.')

def handle_file(cid,uid,prompt):
    answer,err=chat(prompt,uid,'Create the requested content in a clean, file-ready format. Do not wrap it in unnecessary commentary.')
    if err: send_message(cid,'❌ '+err,menu([[btn('🏠 خانه','home')]])); return
    data,name=make_file(prompt,answer)
    if send_document(cid,data,name,'📎 فایل ساخته‌شده توسط Zyro AI'):
        set_mode(uid,'chat')
        send_message(cid,'✅ فایل آماده شد.',menu([[btn('📄 ساخت فایل دیگر','file'),btn('💬 ادامه چت','chat')],[btn('🏠 خانه','home')]]))
    else: send_message(cid,'❌ فایل ساخته شد اما ارسال آن ناموفق بود.')

def handle_tool(cid,uid,action,text):
    prompts={
      'summary':'این متن را دقیق و کوتاه خلاصه کن؛ نکات کلیدی را بولت‌وار بده.',
      'translate':'متن زیر را ترجمه کن و فقط ترجمه طبیعی و دقیق را بده.',
      'code':'بهترین راه‌حل کدنویسی برای درخواست زیر را بده و کد کامل و قابل اجرا ارائه کن.',
      'report':'از درخواست زیر یک گزارش حرفه‌ای با عنوان، خلاصه، بخش‌بندی و جمع‌بندی بساز.'
    }
    handle_text(cid,uid,text,prompts.get(action,''))

def handle_callback(data):
    q=parse_callback(data)
    if not q:return
    cid,uid,action,qid=q; answer_callback(qid)
    if action=='home': set_mode(uid,'chat'); home(cid)
    elif action=='chat': set_mode(uid,'chat'); send_message(cid,'💬 **گفت‌وگوی هوشمند فعال شد**\n\nپیامت را بفرست.',menu([[btn('🖼️ ساخت تصویر','image'),btn('📄 ساخت فایل','file')],[btn('🧠 ابزارهای بیشتر','tools'),btn('🏠 خانه','home')]]))
    elif action=='image': set_mode(uid,'image'); send_message(cid,'🖼️ **حالت ساخت تصویر فعال شد**\n\nتوضیح تصویر را بفرست.',menu([[btn('💬 گفت‌وگو','chat'),btn('🏠 خانه','home')]]))
    elif action=='file': set_mode(uid,'file'); send_message(cid,'📄 **ساخت فایل فعال شد**\n\nمثلاً بنویس: «یک گزارش درباره هوش مصنوعی بساز و PDF بده»',menu([[btn('💬 گفت‌وگو','chat'),btn('🏠 خانه','home')]]))
    elif action in ('summary','translate','code','report'):
        set_mode(uid,action); labels={'summary':'📝 خلاصه‌سازی','translate':'🌍 ترجمه','code':'💻 کدنویسی','report':'📊 گزارش'}
        send_message(cid,f'{labels[action]} فعال شد.\n\nمتنت را بفرست.',menu([[btn('💬 گفت‌وگو','chat'),btn('📄 فایل','file')],[btn('🏠 خانه','home')]]))
    elif action=='tools': tools_menu(cid)
    elif action=='clear': clear_history(uid); set_mode(uid,'chat'); send_message(cid,'🧹 حافظه گفت‌وگو پاک شد. آماده‌ایم.',menu([[btn('💬 شروع گفت‌وگو','chat'),btn('🏠 خانه','home')]]))
    elif action=='help': help_text(cid)

def handle_message(data):
    x=parse_message(data)
    if not x or x['chat_id'] is None:return
    cid,uid,text,typ=x['chat_id'],x['user_id'],x['text'],x['chat_type']
    cmd=text.lower().strip()
    if cmd in ('/start','/ai','ai','zyro'):
        set_mode(uid,'chat'); home(cid); return
    if cmd in ('/clear','/new'):
        clear_history(uid); set_mode(uid,'chat'); send_message(cid,'🧹 گفت‌وگوی جدید آماده است.',menu([[btn('💬 شروع','chat'),btn('🏠 خانه','home')]])); return
    if cmd in ('/help','راهنما'):
        help_text(cid); return
    if typ in ('group','supergroup','channel') and is_mention(text):
        prompt=strip_mention(text)
        if prompt: handle_text(cid,uid,prompt)
        return
    if typ not in ('group','supergroup','channel') and text and not text.startswith('/'):
        mode=get_mode(uid)
        if mode=='image': handle_image(cid,text)
        elif mode=='file': handle_file(cid,uid,text)
        elif mode in ('summary','translate','code','report'): handle_tool(cid,uid,mode,text)
        else: handle_text(cid,uid,text)

# ---------- web ----------
@app.get('/')
def root(): return 'Zyro AI 4.0 is running.'
@app.get('/health')
def health():
    return jsonify(ok=True,zyro=True,version='4.0',chat_model=OPENROUTER_MODEL,image_models=HF_IMAGE_MODELS,image_provider=HF_IMAGE_PROVIDER,features=['chat','image','files','tools','groups','fast-webhook'])

def background(data):
    try:
        handle_callback(data); handle_message(data)
    except Exception:
        traceback.print_exc()

@app.post('/webhook')
def webhook():
    data=request.get_json(silent=True) or {}
    # Ack Telegram/Zento immediately; heavy AI work happens in a background worker.
    threading.Thread(target=background,args=(data,),daemon=True).start()
    return jsonify(ok=True,queued=True)

init_db()
if __name__=='__main__': app.run(host='0.0.0.0',port=int(os.getenv('PORT','8080')),threaded=True)
